from collections import Counter
from datetime import datetime
from io import BytesIO

from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook

from app.db import load_data, next_id, save_data
from app.schemas.employees import EmployeeBulkCreate, EmployeeCreate, EmployeeResponse
from app.services.exports import excel_response, pdf_response


router = APIRouter(prefix="/employees", tags=["employees"])


EMPLOYEE_TEMPLATE_COLUMNS = [
    "code",
    "first_name",
    "last_name",
    "identity_document",
    "birth_date",
    "hire_date",
    "email",
    "phone",
    "address",
    "department",
    "position",
    "contract_type",
    "payment_type",
    "base_salary",
]


def employee_usage_flags(data: dict, employee_id: int) -> tuple[bool, bool]:
    has_salary_advances = any(item["employee_id"] == employee_id for item in data["salary_advances"])
    has_payroll_history = any(item["employee_id"] == employee_id for item in data["payroll_run_details"])
    return has_salary_advances, has_payroll_history


def enrich_employee(item: dict, data: dict) -> dict:
    has_salary_advances, has_payroll_history = employee_usage_flags(data, item["id"])
    return {
        **item,
        "can_delete": not has_salary_advances and not has_payroll_history,
        "has_salary_advances": has_salary_advances,
        "has_payroll_history": has_payroll_history,
    }


@router.get("", response_model=list[EmployeeResponse])
def list_employees() -> list[dict]:
    data = load_data()
    rows = [enrich_employee(item, data) for item in data["employees"]]
    return sorted(rows, key=lambda item: item["id"], reverse=True)


@router.get("/export.xlsx")
def export_employees_excel() -> StreamingResponse:
    data = load_data()
    rows = [
        [
            item.get("code"),
            f"{item.get('first_name', '')} {item.get('last_name', '')}".strip(),
            item.get("identity_document"),
            item.get("hire_date"),
            item.get("department"),
            item.get("position"),
            item.get("payment_type"),
            item.get("base_salary"),
        ]
        for item in data["employees"]
    ]
    return excel_response(
        "employees_repository.xlsx",
        "employees",
        ["Codigo", "Empleado", "Documento", "Ingreso", "Departamento", "Cargo", "Ciclo", "Salario"],
        rows,
    )


@router.get("/export.pdf")
def export_employees_pdf() -> StreamingResponse:
    data = load_data()
    rows = [
        [
            item.get("code"),
            f"{item.get('first_name', '')} {item.get('last_name', '')}".strip(),
            item.get("identity_document") or "",
            item.get("hire_date") or "",
            item.get("department") or "",
            item.get("position") or "",
            item.get("payment_type") or "",
            item.get("base_salary") or 0,
        ]
        for item in data["employees"]
    ]
    return pdf_response(
        "employees_repository.pdf",
        "Repositorio de Empleados",
        ["Codigo", "Empleado", "Documento", "Ingreso", "Departamento", "Cargo", "Ciclo", "Salario"],
        rows,
    )


@router.get("/template")
def employee_template() -> StreamingResponse:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "employees"
    sheet.append(EMPLOYEE_TEMPLATE_COLUMNS)
    sheet.append(
        [
            "EMP-401",
            "Laura",
            "Santos",
            "001-000000-0001A",
            "1998-03-15",
            "2026-04-07",
            "laura@demo.com",
            "8888-1111",
            "Managua",
            "Finanzas",
            "Analista",
            "permanent",
            "monthly",
            14500,
        ]
    )

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    headers = {
        "Content-Disposition": 'attachment; filename="employee_import_template.xlsx"'
    }
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


@router.post("", response_model=EmployeeResponse, status_code=201)
def create_employee(payload: EmployeeCreate) -> dict:
    data = load_data()
    exists = next(
        (
            item
            for item in data["employees"]
            if item["company_id"] == payload.company_id and item["code"] == payload.code
        ),
        None,
    )
    if exists is not None:
        raise HTTPException(status_code=409, detail="Ya existe un empleado con ese codigo")

    employee = payload.model_dump()
    employee["id"] = next_id(data["employees"])
    data["employees"].append(employee)
    save_data(data)
    return employee


@router.post("/bulk", status_code=201)
def bulk_create_employees(payload: EmployeeBulkCreate) -> dict:
    data = load_data()
    created = []
    skipped = []

    for raw_employee in payload.employees:
        employee_data = raw_employee.model_dump()
        employee_data["company_id"] = payload.company_id
        exists = next(
            (
                item
                for item in data["employees"]
                if item["company_id"] == payload.company_id and item["code"] == employee_data["code"]
            ),
            None,
        )
        if exists is not None:
            skipped.append(
                {
                    "code": employee_data["code"],
                    "reason": "Codigo duplicado",
                }
            )
            continue

        employee_data["id"] = next_id(data["employees"])
        data["employees"].append(employee_data)
        created.append(employee_data)

    save_data(data)
    return {
        "created_count": len(created),
        "skipped_count": len(skipped),
        "created": created,
        "skipped": skipped,
    }


@router.post("/bulk-excel", status_code=201)
async def bulk_create_employees_excel(file: UploadFile = File(...), company_id: int = 1) -> dict:
    if not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Solo se admite archivo .xlsx")

    content = await file.read()
    workbook = load_workbook(filename=BytesIO(content))
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="El archivo Excel esta vacio")

    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    missing_headers = [column for column in EMPLOYEE_TEMPLATE_COLUMNS if column not in headers]
    if missing_headers:
        raise HTTPException(
            status_code=400,
            detail=f"Faltan columnas requeridas: {', '.join(missing_headers)}",
        )

    employees_payload = []
    for row in rows[1:]:
        if not row or all(value in (None, "") for value in row):
            continue
        raw = dict(zip(headers, row))
        employees_payload.append(
            EmployeeCreate(
                company_id=company_id,
                code=str(raw.get("code", "")).strip(),
                first_name=str(raw.get("first_name", "")).strip(),
                last_name=str(raw.get("last_name", "")).strip(),
                identity_document=None if raw.get("identity_document") in (None, "") else str(raw.get("identity_document")).strip(),
                birth_date=None if raw.get("birth_date") in (None, "") else str(raw.get("birth_date"))[:10],
                hire_date=None if raw.get("hire_date") in (None, "") else str(raw.get("hire_date"))[:10],
                email=None if raw.get("email") in (None, "") else str(raw.get("email")).strip(),
                phone=None if raw.get("phone") in (None, "") else str(raw.get("phone")).strip(),
                address=None if raw.get("address") in (None, "") else str(raw.get("address")).strip(),
                department=None if raw.get("department") in (None, "") else str(raw.get("department")).strip(),
                position=str(raw.get("position", "")).strip(),
                contract_type=str(raw.get("contract_type", "permanent")).strip() or "permanent",
                payment_type=str(raw.get("payment_type", "monthly")).strip() or "monthly",
                base_salary=float(raw.get("base_salary") or 0),
                active=True,
            )
        )

    return bulk_create_employees(EmployeeBulkCreate(company_id=company_id, employees=employees_payload))


@router.get("/reports/hiring")
def hiring_report(group_by: str = "monthly") -> dict:
    data = load_data()
    if group_by not in {"monthly", "biweekly", "annual"}:
        raise HTTPException(status_code=400, detail="El reporte solo soporta monthly, biweekly o annual")

    buckets = Counter()
    for employee in data["employees"]:
        hire_date = employee.get("hire_date")
        if not hire_date:
            continue

        dt = datetime.strptime(hire_date, "%Y-%m-%d")
        if group_by == "monthly":
            bucket = dt.strftime("%Y-%m")
        elif group_by == "annual":
            bucket = dt.strftime("%Y")
        else:
            half = "Q1" if dt.day <= 15 else "Q2"
            bucket = f"{dt.strftime('%Y-%m')}-{half}"
        buckets[bucket] += 1

    rows = [{"period": key, "count": buckets[key]} for key in sorted(buckets.keys())]
    return {
        "group_by": group_by,
        "rows": rows,
        "total_employees_with_hire_date": sum(item["count"] for item in rows),
    }


@router.get("/reports/hiring/export.xlsx")
def export_hiring_report_excel(group_by: str = "monthly") -> StreamingResponse:
    report = hiring_report(group_by)
    rows = [[item["period"], item["count"]] for item in report["rows"]]
    return excel_response("employees_hiring_report.xlsx", "hiring-report", ["Periodo", "Contratados"], rows)


@router.get("/reports/hiring/export.pdf")
def export_hiring_report_pdf(group_by: str = "monthly") -> StreamingResponse:
    report = hiring_report(group_by)
    rows = [[item["period"], item["count"]] for item in report["rows"]]
    return pdf_response("employees_hiring_report.pdf", "Reporte de Contrataciones", ["Periodo", "Contratados"], rows)


@router.get("/{employee_id}", response_model=EmployeeResponse)
def get_employee(employee_id: int) -> dict:
    data = load_data()
    row = next((item for item in data["employees"] if item["id"] == employee_id), None)
    if row is None:
        raise HTTPException(status_code=404, detail="Empleado no encontrado")
    return enrich_employee(row, data)


@router.patch("/{employee_id}/deactivate")
def deactivate_employee(employee_id: int) -> dict:
    data = load_data()
    employee = next((item for item in data["employees"] if item["id"] == employee_id), None)
    if employee is None:
        raise HTTPException(status_code=404, detail="Empleado no encontrado")

    employee["active"] = False
    employee["employment_status"] = "inactive"
    save_data(data)
    return {"message": "Empleado desactivado"}


@router.patch("/{employee_id}/archive")
def archive_employee(employee_id: int) -> dict:
    data = load_data()
    employee = next((item for item in data["employees"] if item["id"] == employee_id), None)
    if employee is None:
        raise HTTPException(status_code=404, detail="Empleado no encontrado")

    employee["active"] = False
    employee["employment_status"] = "archived"
    save_data(data)
    return {"message": "Empleado archivado"}


@router.delete("/{employee_id}")
def delete_employee(employee_id: int) -> dict:
    data = load_data()
    employee = next((item for item in data["employees"] if item["id"] == employee_id), None)
    if employee is None:
        raise HTTPException(status_code=404, detail="Empleado no encontrado")

    has_salary_advances, has_payroll_history = employee_usage_flags(data, employee_id)
    if has_salary_advances or has_payroll_history:
        raise HTTPException(
            status_code=409,
            detail="No se puede eliminar este empleado porque tiene adelantos o historial de planilla",
        )

    data["employees"] = [item for item in data["employees"] if item["id"] != employee_id]
    save_data(data)
    return {"message": "Empleado eliminado"}
