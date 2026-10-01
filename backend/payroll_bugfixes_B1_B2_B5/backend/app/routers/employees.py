from collections import Counter
from datetime import date, datetime
from io import BytesIO

from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from app.db import load_data, next_id, save_data
from app.schemas.employees import EmployeeBulkCreate, EmployeeCreate, EmployeeResponse
from app.services.exports import excel_response, pdf_response


router = APIRouter(prefix="/employees", tags=["employees"])


EMPLOYEE_TEMPLATE_COLUMNS = [
    "code",
    "first_name",
    "last_name",
    "identity_document",
    "birth_date",
    "hire_date",
    "email",
    "phone",
    "address",
    "department",
    "position",
    "contract_type",
    "payment_type",
    "base_salary",
]


def employee_usage_flags(data: dict, employee_id: int) -> tuple[bool, bool]:
    has_salary_advances = any(item["employee_id"] == employee_id for item in data["salary_advances"])
    has_payroll_history = any(item["employee_id"] == employee_id for item in data["payroll_run_details"])
    return has_salary_advances, has_payroll_history


def enrich_employee(item: dict, data: dict) -> dict:
    has_salary_advances, has_payroll_history = employee_usage_flags(data, item["id"])
    return {
        **item,
        "can_delete": not has_salary_advances and not has_payroll_history,
        "has_salary_advances": has_salary_advances,
        "has_payroll_history": has_payroll_history,
    }


@router.get("", response_model=list[EmployeeResponse])
def list_employees() -> list[dict]:
    data = load_data()
    rows = [enrich_employee(item, data) for item in data["employees"]]
    return sorted(rows, key=lambda item: item["id"], reverse=True)


@router.get("/export.xlsx")
def export_employees_excel() -> StreamingResponse:
    data = load_data()
    rows = [
        [
            item.get("code"),
            f"{item.get('first_name', '')} {item.get('last_name', '')}".strip(),
            item.get("identity_document"),
            item.get("hire_date"),
            item.get("department"),
            item.get("position"),
            item.get("payment_type"),
            item.get("base_salary"),
        ]
        for item in data["employees"]
    ]
    return excel_response(
        "employees_repository.xlsx",
        "employees",
        ["Codigo", "Empleado", "Documento", "Ingreso", "Departamento", "Cargo", "Ciclo", "Salario"],
        rows,
    )


@router.get("/export.pdf")
def export_employees_pdf() -> StreamingResponse:
    data = load_data()
    rows = [
        [
            item.get("code"),
            f"{item.get('first_name', '')} {item.get('last_name', '')}".strip(),
            item.get("identity_document") or "",
            item.get("hire_date") or "",
            item.get("department") or "",
            item.get("position") or "",
            item.get("payment_type") or "",
            item.get("base_salary") or 0,
        ]
        for item in data["employees"]
    ]
    return pdf_response(
        "employees_repository.pdf",
        "Repositorio de Empleados",
        ["Codigo", "Empleado", "Documento", "Ingreso", "Departamento", "Cargo", "Ciclo", "Salario"],
        rows,
    )


@router.get("/template")
def employee_template() -> StreamingResponse:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "employees"
    sheet.append(EMPLOYEE_TEMPLATE_COLUMNS)
    sheet.append(
        [
            "EMP-401",
            "Laura",
            "Santos",
            "001-000000-0001A",
            "1998-03-15",
            "2026-04-07",
            "laura@demo.com",
            "8888-1111",
            "Managua",
            "Finanzas",
            "Analista",
            "permanent",
            "monthly",
            14500,
        ]
    )

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    headers = {
        "Content-Disposition": 'attachment; filename="employee_import_template.xlsx"'
    }
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


@router.post("", response_model=EmployeeResponse, status_code=201)
def create_employee(payload: EmployeeCreate) -> dict:
    data = load_data()
    exists = next(
        (
            item
            for item in data["employees"]
            if item["company_id"] == payload.company_id and item["code"] == payload.code
        ),
        None,
    )
    if exists is not None:
        raise HTTPException(status_code=409, detail="Ya existe un empleado con ese codigo")

    employee = payload.model_dump()
    employee["id"] = next_id(data["employees"])
    data["employees"].append(employee)
    save_data(data)
    return employee


@router.post("/bulk", status_code=201)
def bulk_create_employees(payload: EmployeeBulkCreate) -> dict:
    data = load_data()
    created = []
    skipped = []

    for raw_employee in payload.employees:
        employee_data = raw_employee.model_dump()
        employee_data["company_id"] = payload.company_id
        exists = next(
            (
                item
                for item in data["employees"]
                if item["company_id"] == payload.company_id and item["code"] == employee_data["code"]
            ),
            None,
        )
        if exists is not None:
            skipped.append(
                {
                    "code": employee_data["code"],
                    "reason": "Codigo duplicado",
                }
            )
            continue

        employee_data["id"] = next_id(data["employees"])
        data["employees"].append(employee_data)
        created.append(employee_data)

    save_data(data)
    return {
        "created_count": len(created),
        "skipped_count": len(skipped),
        "created": created,
        "skipped": skipped,
    }


def _excel_text(value) -> str | None:
    if value in (None, ""):
        return None
    text = str(value).strip()
    return text or None


def _excel_date(value, field: str) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()[:10]
    try:
        datetime.strptime(text, "%Y-%m-%d")
    except ValueError:
        raise ValueError(f"{field}: fecha invalida '{value}', use el formato AAAA-MM-DD") from None
    return text


@router.post("/bulk-excel", status_code=201)
async def bulk_create_employees_excel(file: UploadFile = File(...), company_id: int = 1) -> dict:
    if not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Solo se admite archivo .xlsx")

    content = await file.read()
    try:
        workbook = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
    except Exception:
        # B6: un archivo danado o que no es Excel es un error del usuario, no del servidor
        raise HTTPException(status_code=400, detail="El archivo no es un Excel .xlsx valido")

    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="El archivo Excel esta vacio")

    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    missing_headers = [column for column in EMPLOYEE_TEMPLATE_COLUMNS if column not in headers]
    if missing_headers:
        raise HTTPException(
            status_code=400,
            detail=f"Faltan columnas requeridas: {', '.join(missing_headers)}",
        )

    # B5: cada fila se valida por separado; una fila invalida ya no detiene la importacion
    employees_payload = []
    invalid_rows = []
    for row_number, row in enumerate(rows[1:], start=2):
        if not row or all(value in (None, "") for value in row):
            continue
        raw = dict(zip(headers, row))
        code = _excel_text(raw.get("code")) or ""
        try:
            employees_payload.append(
                EmployeeCreate(
                    company_id=company_id,
                    code=code,
                    first_name=_excel_text(raw.get("first_name")) or "",
                    last_name=_excel_text(raw.get("last_name")) or "",
                    identity_document=_excel_text(raw.get("identity_document")),
                    birth_date=_excel_date(raw.get("birth_date"), "birth_date"),
                    hire_date=_excel_date(raw.get("hire_date"), "hire_date"),
                    email=_excel_text(raw.get("email")),
                    phone=_excel_text(raw.get("phone")),
                    address=_excel_text(raw.get("address")),
                    department=_excel_text(raw.get("department")),
                    position=_excel_text(raw.get("position")) or "",
                    contract_type=_excel_text(raw.get("contract_type")) or "permanent",
                    payment_type=_excel_text(raw.get("payment_type")) or "monthly",
                    base_salary=raw.get("base_salary"),
                    active=True,
                )
            )
        except ValidationError as error:
            messages = [f"{'.'.join(str(part) for part in item['loc'])}: {item['msg']}" for item in error.errors()]
            invalid_rows.append({"row": row_number, "code": code, "reason": "Fila invalida", "errors": messages})
        except ValueError as error:
            invalid_rows.append({"row": row_number, "code": code, "reason": "Fila invalida", "errors": [str(error)]})

    result = bulk_create_employees(EmployeeBulkCreate(company_id=company_id, employees=employees_payload))
    result["skipped"] = invalid_rows + result["skipped"]
    result["skipped_count"] = len(result["skipped"])
    result["invalid_count"] = len(invalid_rows)
    return result


@router.get("/reports/hiring")
def hiring_report(group_by: str = "monthly") -> dict:
    data = load_data()
    if group_by not in {"monthly", "biweekly", "annual"}:
        raise HTTPException(status_code=400, detail="El reporte solo soporta monthly, biweekly o annual")

    buckets = Counter()
    for employee in data["employees"]:
        hire_date = employee.get("hire_date")
        if not hire_date:
            continue

        dt = datetime.strptime(hire_date, "%Y-%m-%d")
        if group_by == "monthly":
            bucket = dt.strftime("%Y-%m")
        elif group_by == "annual":
            bucket = dt.strftime("%Y")
        else:
            half = "Q1" if dt.day <= 15 else "Q2"
            bucket = f"{dt.strftime('%Y-%m')}-{half}"
        buckets[bucket] += 1

    rows = [{"period": key, "count": buckets[key]} for key in sorted(buckets.keys())]
    return {
        "group_by": group_by,
        "rows": rows,
        "total_employees_with_hire_date": sum(item["count"] for item in rows),
    }


@router.get("/reports/hiring/export.xlsx")
def export_hiring_report_excel(group_by: str = "monthly") -> StreamingResponse:
    report = hiring_report(group_by)
    rows = [[item["period"], item["count"]] for item in report["rows"]]
    return excel_response("employees_hiring_report.xlsx", "hiring-report", ["Periodo", "Contratados"], rows)


@router.get("/reports/hiring/export.pdf")
def export_hiring_report_pdf(group_by: str = "monthly") -> StreamingResponse:
    report = hiring_report(group_by)
    rows = [[item["period"], item["count"]] for item in report["rows"]]
    return pdf_response("employees_hiring_report.pdf", "Reporte de Contrataciones", ["Periodo", "Contratados"], rows)


@router.get("/{employee_id}", response_model=EmployeeResponse)
def get_employee(employee_id: int) -> dict:
    data = load_data()
    row = next((item for item in data["employees"] if item["id"] == employee_id), None)
    if row is None:
        raise HTTPException(status_code=404, detail="Empleado no encontrado")
    return enrich_employee(row, data)


@router.patch("/{employee_id}/deactivate")
def deactivate_employee(employee_id: int) -> dict:
    data = load_data()
    employee = next((item for item in data["employees"] if item["id"] == employee_id), None)
    if employee is None:
        raise HTTPException(status_code=404, detail="Empleado no encontrado")

    employee["active"] = False
    employee["employment_status"] = "inactive"
    save_data(data)
    return {"message": "Empleado desactivado"}


@router.patch("/{employee_id}/archive")
def archive_employee(employee_id: int) -> dict:
    data = load_data()
    employee = next((item for item in data["employees"] if item["id"] == employee_id), None)
    if employee is None:
        raise HTTPException(status_code=404, detail="Empleado no encontrado")

    employee["active"] = False
    employee["employment_status"] = "archived"
    save_data(data)
    return {"message": "Empleado archivado"}


@router.delete("/{employee_id}")
def delete_employee(employee_id: int) -> dict:
    data = load_data()
    employee = next((item for item in data["employees"] if item["id"] == employee_id), None)
    if employee is None:
        raise HTTPException(status_code=404, detail="Empleado no encontrado")

    has_salary_advances, has_payroll_history = employee_usage_flags(data, employee_id)
    if has_salary_advances or has_payroll_history:
        raise HTTPException(
            status_code=409,
            detail="No se puede eliminar este empleado porque tiene adelantos o historial de planilla",
        )

    data["employees"] = [item for item in data["employees"] if item["id"] != employee_id]
    save_data(data)
    return {"message": "Empleado eliminado"}
